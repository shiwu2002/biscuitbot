"""文档文本抽取工具。

所属模块与项目作用
===================
本文件位于 biscuitbot/utils 目录，是工具函数模块的文档抽取组件。
在项目架构中起到的作用：
将用户附件（PDF / DOCX / XLSX / PPTX / 纯文本 / 图片等）抽取为纯文本，
并把图片与文档分离，使下游 LLM 层只需处理文本块与视觉块。
为节省启动开销，各格式的解析库均采用惰性导入。
"""

import mimetypes  # 扩展名到 MIME 的回退嗅探
from pathlib import Path

from loguru import logger  # 结构化日志记录

from biscuitbot.utils.helpers import detect_image_mime  # 基于魔数的图片 MIME 嗅探

# 支持文本抽取的文件扩展名集合
SUPPORTED_EXTENSIONS: set[str] = {
    # Document formats
    ".pdf",
    ".docx",
    ".xlsx",
    ".pptx",
    # Text formats
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".xml",
    ".html",
    ".htm",
    ".log",
    ".yaml",
    ".yml",
    ".toml",
    ".ini",
    ".cfg",
    # Image formats (for future OCR support)
    ".png",
    ".jpg",
    ".jpeg",
    ".gif",
    ".webp",
}

# 单文件抽取文本的最大字符数，超出将被截断
_MAX_TEXT_LENGTH = 200_000


def extract_text(path: Path) -> str | None:
    """从文件中抽取文本。

    Args:
        path: 文件路径。

    Returns:
        抽取出的文本字符串；不支持的扩展名返回 None；
        抽取失败返回以 ``[error: ...]`` 开头的错误字符串。
    """
    if not isinstance(path, Path):
        path = Path(path)

    if not path.exists():
        return f"[error: file not found: {path}]"

    ext = path.suffix.lower()

    # Document formats -- each branch lazily imports its parser so that
    # startup does not pay the ~25 MB cost of loading openpyxl /
    # python-docx / python-pptx / pypdf up front (see issue #3422).
    if ext == ".pdf":
        return _extract_pdf(path)
    elif ext == ".docx":
        return _extract_docx(path)
    elif ext == ".xlsx":
        return _extract_xlsx(path)
    elif ext == ".pptx":
        return _extract_pptx(path)
    elif _is_text_extension(ext):
        return _extract_text_file(path)
    elif ext in {".png", ".jpg", ".jpeg", ".gif", ".webp"}:
        # Image files - for future OCR support
        return f"[image: {path.name}]"
    else:
        # Unsupported extension
        return None


def _extract_pdf(path: Path) -> str:
    """使用 pypdf 从 PDF 中抽取文本。"""
    try:
        from pypdf import PdfReader  # 惰性导入 PDF 解析库
    except ImportError:
        return "[error: pypdf not installed]"
    try:
        reader = PdfReader(path)
        pages: list[str] = []
        for i, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            pages.append(f"--- Page {i} ---\n{text}")
        return _truncate("\n\n".join(pages), _MAX_TEXT_LENGTH)
    except Exception as e:
        logger.exception("Failed to extract PDF {}", path)
        return f"[error: failed to extract PDF: {e!s}]"


def _extract_docx(path: Path) -> str:
    """使用 python-docx 从 DOCX 中抽取文本。"""
    try:
        from docx import Document as DocxDocument  # 惰性导入 DOCX 解析库
    except ImportError:
        return "[error: python-docx not installed]"
    try:
        doc = DocxDocument(path)
        paragraphs: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]
        return _truncate("\n\n".join(paragraphs), _MAX_TEXT_LENGTH)
    except Exception as e:
        logger.exception("Failed to extract DOCX {}", path)
        return f"[error: failed to extract DOCX: {e!s}]"


def _extract_xlsx(path: Path) -> str:
    """使用 openpyxl 从 XLSX 中抽取文本。"""
    try:
        from openpyxl import load_workbook  # 惰性导入表格解析库
    except ImportError:
        return "[error: openpyxl not installed]"
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            sheets: list[str] = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows: list[str] = []
                for row in ws.iter_rows(values_only=True):
                    # 以制表符分隔单元格，空单元格留空
                    row_text = "\t".join(str(cell) if cell is not None else "" for cell in row)
                    if row_text.strip():
                        rows.append(row_text)
                if rows:
                    sheets.append(f"--- Sheet: {sheet_name} ---\n" + "\n".join(rows))
            return _truncate("\n\n".join(sheets), _MAX_TEXT_LENGTH)
        finally:
            wb.close()
    except Exception as e:
        logger.exception("Failed to extract XLSX {}", path)
        return f"[error: failed to extract XLSX: {e!s}]"


def _extract_pptx(path: Path) -> str:
    """使用 python-pptx 从 PPTX 中抽取文本。"""
    try:
        from pptx import Presentation as PptxPresentation  # 惰性导入 PPT 解析库
    except ImportError:
        return "[error: python-pptx not installed]"
    try:
        prs = PptxPresentation(path)
        slides: list[str] = []
        for i, slide in enumerate(prs.slides, 1):
            slide_text: list[str] = []
            for shape in slide.shapes:
                _collect_pptx_shape_text(shape, slide_text)
            if slide_text:
                slides.append(f"--- Slide {i} ---\n" + "\n".join(slide_text))
        return _truncate("\n\n".join(slides), _MAX_TEXT_LENGTH)
    except Exception as e:
        logger.exception("Failed to extract PPTX {}", path)
        return f"[error: failed to extract PPTX: {e!s}]"


def _collect_pptx_shape_text(shape, out: list[str]) -> None:
    """收集 PPTX 形状中的文本，递归处理分组与表格。

    分组形状 ``has_text_frame=False``，必须通过 ``.shapes`` 遍历；
    表格是 GraphicFrame 对象，其单元格文本位于 ``.table`` 之下。
    """
    sub_shapes = getattr(shape, "shapes", None)  # 分组形状：递归子形状
    if sub_shapes is not None:
        for sub in sub_shapes:
            _collect_pptx_shape_text(sub, out)
        return

    if getattr(shape, "has_table", False):  # 表格：按行收集单元格文本
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = "\t".join(cell for cell in cells if cell)
            if line:
                out.append(line)
        return

    text = getattr(shape, "text", "")  # 普通形状：直接取 text
    if text:
        out.append(text)


def _extract_text_file(path: Path) -> str:
    """从纯文本文件中抽取文本。"""
    try:
        # Try UTF-8 first, then latin-1 fallback
        try:
            content = path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            content = path.read_text(encoding="latin-1")  # 兜底编码，避免解码失败
        return _truncate(content, _MAX_TEXT_LENGTH)
    except Exception as e:
        logger.exception("Failed to read text file {}", path)
        return f"[error: failed to read file: {e!s}]"


def _truncate(text: str, max_length: int) -> str:
    """截断文本，并在末尾追加截断标记。"""
    if len(text) <= max_length:
        return text
    return text[:max_length] + f"... (truncated, {len(text)} chars total)"


def _is_text_extension(ext: str) -> bool:
    """判断扩展名是否属于纯文本格式。"""
    return ext in {
        ".txt",
        ".md",
        ".csv",
        ".json",
        ".xml",
        ".html",
        ".htm",
        ".log",
        ".yaml",
        ".yml",
        ".toml",
        ".ini",
        ".cfg",
    }


# ---------------------------------------------------------------------------
# High-level helper: split media into images + extracted document text
# ---------------------------------------------------------------------------

# 抽取单文件的最大字节数（50 MB），超出则跳过以限制内存/CPU
_MAX_EXTRACT_FILE_SIZE = 50 * 1024 * 1024  # 50 MB


def is_image_file(path: str) -> bool:
    """判断 *path* 是否为图片文件。

    优先用魔数检测（读取前 16 字节），失败时回退到 ``mimetypes``
    基于扩展名的判断。
    """
    p = Path(path)
    mime: str | None = None
    if p.is_file():
        try:
            with p.open("rb") as f:
                mime = detect_image_mime(f.read(16))  # 读取前 16 字节做魔数嗅探
        except OSError:
            mime = None
    if not mime:
        mime = mimetypes.guess_type(path)[0]  # 回退到扩展名判断
    return bool(mime and mime.startswith("image/"))


def reference_non_image_attachments(
    content: str, media: list[str],
) -> tuple[str, list[str]]:
    """在不读取文件内容的前提下，将图片与非图片附件分离。

    图片路径保留，供下游视觉块构造使用；
    非图片路径以 ``[Attachment: path]`` 形式追加到内容末尾。
    """
    image_paths: list[str] = []
    attachment_refs: list[str] = []
    for path in media:
        if is_image_file(path):
            image_paths.append(path)
        else:
            attachment_refs.append(f"[Attachment: {path}]")
    if attachment_refs:
        suffix = "\n".join(attachment_refs)
        content = f"{content}\n\n{suffix}" if content else suffix
    return content, image_paths


def extract_documents(
    text: str,
    media_paths: list[str],
    *,
    max_file_size: int = _MAX_EXTRACT_FILE_SIZE,
) -> tuple[str, list[str]]:
    """将 *media_paths* 中的图片与文档分离。

    文档（PDF、DOCX、XLSX、PPTX、纯文本等）会被抽取为文本并
    追加到 *text* 中；返回列表只保留图片路径，使下游层只需处理
    视觉块。

    超过 *max_file_size* 字节的文件会被跳过并记录警告，以避免
    无限制的内存 / CPU 占用。
    """
    image_paths: list[str] = []
    doc_texts: list[str] = []

    for path_str in media_paths:
        p = Path(path_str)
        if not p.is_file():
            continue

        try:
            size = p.stat().st_size
        except OSError:
            continue
        if size > max_file_size:  # 超限文件跳过
            logger.warning(
                "Skipping oversized file for extraction: {} ({:.1f} MB > {} MB limit)",
                p.name, size / (1024 * 1024), max_file_size // (1024 * 1024),
            )
            continue

        if is_image_file(path_str):
            image_paths.append(path_str)
        else:
            extracted = extract_text(p)
            if extracted and not extracted.startswith("[error:"):
                doc_texts.append(f"[File: {p.name}]\n{extracted}")

    if doc_texts:  # 将抽取的文档文本拼接到原文本后
        text = text + "\n\n" + "\n\n".join(doc_texts)

    return text, image_paths
